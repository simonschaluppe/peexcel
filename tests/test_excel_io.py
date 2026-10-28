"""Independent workbook fixtures: never use the pexl writer to test its reader."""

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from pexl import Project
from pexl.io.excel import build_project_in_dataframe, write_project_excel


@pytest.fixture
def workbook(tmp_path):
    def write(name="project.xlsx", *, sheets):
        path = tmp_path / name
        path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        book.remove(book.active)
        for sheet_name, rows in sheets.items():
            sheet = book.create_sheet(sheet_name)
            for row in rows:
                sheet.append(row)
        book.save(path)
        book.close()
        return path

    return write


def annual_sheets():
    return {
        "IN": [
            ["var_name", "Name", "Demo | B", "Demo | A"],
            ["project_name", "Project", "Demo", "Demo"],
            ["project_scenario_name", "Scenario", "B", "A"],
            ["preset_recorded_heating_system", "Heating", "Wärmepumpe", None],
        ],
        # Different column order: values must join by identity.
        "OUT": [["var_name", "Demo | A", "Demo | B"], ["QH", 0, 42.5]],
    }


def timestep_sheets(value=10, out_column="Live"):
    return {
        "IN": [["var_name", "STAGE"], ["project_name", "Demo"]],
        "OUT": [["var_name", out_column], ["QH", value]],
        "SIM2": [
            ["date", "Ta"],
            ["2026-01-01 00:00:00", value],
            ["2026-01-01 01:00:00", value + 1],
        ],
    }


def test_annual_import_preserves_identity_types_and_missing_values(workbook):
    path = workbook(sheets=annual_sheets())
    project = Project.from_excel(path, unknown="raise")
    assert project.column_names() == ["Demo | B", "Demo | A"]
    assert project["Demo | B"].to_var_dict()["QH"] == 42.5
    assert project["Demo | A"].to_var_dict()["QH"] == 0
    assert project[0].to_var_dict()["preset_recorded_heating_system"] == "Wärmepumpe"
    assert pd.isna(project[1].to_var_dict()["preset_recorded_heating_system"])
    assert not project[0].timeseries.available
    assert project.file_source == path
    assert project.warnings == []


def test_sim2_uses_filename_and_loads_timestamps(workbook):
    project = Project.from_excel(workbook("AT001.xlsx", sheets=timestep_sheets()))
    assert project.column_names() == ["AT001"]
    assert project[0].to_var_dict()["QH"] == 10
    frame = project[0].timeseries.to_frame()
    assert frame["Ta"].tolist() == [10, 11]
    assert frame.index.tolist() == list(pd.date_range("2026-01-01", periods=2, freq="h"))
    assert project.warnings == []


def test_sim2_fallback_out_column_warns(workbook):
    project = Project.from_excel(workbook(sheets=timestep_sheets(out_column="Result")))
    assert project[0].to_var_dict()["QH"] == 10
    assert any("using 'Result'" in warning for warning in project.warnings)


def test_from_files_preserves_order_timeseries_and_source_warnings(workbook):
    first = workbook("B.xlsx", sheets=timestep_sheets(20, "Result"))
    second = workbook("A.xlsx", sheets=timestep_sheets(30))
    project = Project.from_files([first, second], unknown="raise")
    assert project.column_names() == ["B", "A"]
    assert [s.to_var_dict()["QH"] for s in project] == [20, 30]
    assert project[0].timeseries.to_frame()["Ta"].tolist() == [20, 21]
    assert project[1].timeseries.to_frame()["Ta"].tolist() == [30, 31]
    assert any(warning.startswith("B.xlsx:") for warning in project.warnings)


def test_from_files_rejects_duplicate_identities_across_folders(workbook):
    first = workbook("old/AT001.xlsx", sheets=timestep_sheets())
    second = workbook("new/AT001.xlsx", sheets=timestep_sheets(30))
    with pytest.raises(ValueError, match="Duplicate scenario"):
        Project.from_files([first, second])


@pytest.mark.parametrize("missing", ["IN", "OUT"])
def test_missing_required_sheet_is_rejected(workbook, missing):
    sheets = annual_sheets()
    del sheets[missing]
    with pytest.raises(ValueError, match="must contain sheets"):
        Project.from_excel(workbook(sheets=sheets))


def test_sim2_requires_stage_input(workbook):
    sheets = timestep_sheets()
    sheets["IN"][0][1] = "Variant"
    with pytest.raises(ValueError, match="STAGE"):
        Project.from_excel(workbook(sheets=sheets))


def test_unknown_variables_can_be_rejected_or_ignored(workbook):
    sheets = annual_sheets()
    sheets["IN"].append(["not_a_schema_variable", "Unknown", 1, 2])
    path = workbook(sheets=sheets)
    with pytest.raises(KeyError, match="not_a_schema_variable"):
        Project.from_excel(path, unknown="raise")
    project = Project.from_excel(path, unknown="ignore")
    assert project[0].to_var_dict()["QH"] == 42.5
    assert "not_a_schema_variable" not in project[0].to_var_dict()


def test_in_table_preserves_requested_layout_without_mutating_project(workbook):
    project = Project.from_excel(workbook(sheets=annual_sheets()))
    before = [scenario.to_var_dict() for scenario in project]
    columns = ["var_name", "unit"]
    frame = build_project_in_dataframe(project, base_columns=columns)
    assert frame.columns.tolist() == ["var_name", "unit", "Demo | B", "Demo | A"]
    assert columns == ["var_name", "unit"]
    assert not frame["var_name"].duplicated().any()
    assert frame.set_index("var_name").loc["project_name", "Demo | B"] == "Demo"
    assert [scenario.to_var_dict() for scenario in project] == before


def test_empty_in_table_and_invalid_metadata_columns():
    frame = build_project_in_dataframe(Project(), base_columns=["var_name"])
    assert frame.empty
    assert frame.columns.tolist() == ["var_name"]
    with pytest.raises(ValueError, match="must include 'var_name'"):
        build_project_in_dataframe(Project(), base_columns=["unit"])


@pytest.mark.parametrize("include_meta", [True, False])
def test_existing_export_is_input_only_and_preserves_values(workbook, tmp_path, include_meta):
    project = Project.from_excel(workbook(sheets=annual_sheets()))
    target = tmp_path / "inputs.xlsx"
    assert project.to_excel(target, include_meta=include_meta) == target
    book = load_workbook(target, read_only=True, data_only=True)
    try:
        assert book.sheetnames == (["IN", "META"] if include_meta else ["IN"])
        rows = list(book["IN"].values)
        headers = rows[0]
        key_index = headers.index("var_name")
        scenario_index = headers.index("Demo | B")
        assert headers[-2:] == ("Demo | B", "Demo | A")
        assert "default" not in headers
        values = {row[key_index]: row[scenario_index] for row in rows[1:]}
        assert values["project_name"] == "Demo"
        assert values["preset_recorded_heating_system"] == "Wärmepumpe"
    finally:
        book.close()


@pytest.mark.parametrize("include_meta", [True, False])
def test_annual_export_workbook_and_roundtrip(workbook, tmp_path, include_meta):
    project = Project.from_excel(workbook(sheets=annual_sheets()), unknown="raise")
    target = tmp_path / "annual.xlsx"
    assert project.to_excel(target, include_out=True, include_meta=include_meta) == target
    book = load_workbook(target, read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["IN", "OUT"] + (["META"] if include_meta else [])
        for sheet in ("IN", "OUT"):
            rows = list(book[sheet].values)
            assert rows[0][-2:] == ("Demo | B", "Demo | A")
            assert "default" not in rows[0]
            key = rows[0].index("var_name")
            assert len({row[key] for row in rows[1:]}) == len(rows) - 1
            if sheet == "OUT":
                values = {row[key]: row[-2:] for row in rows[1:]}
                assert values["QH"] == (42.5, 0)
    finally:
        book.close()
    restored = Project.from_excel(target, unknown="raise")
    assert restored.column_names() == project.column_names()
    assert restored.warnings == []
    for original, loaded in zip(project, restored):
        for name, value in original.to_var_dict().items():
            actual = loaded.to_var_dict()[name]
            if pd.isna(value):
                assert pd.isna(actual), name
            elif isinstance(value, (int, float)):
                assert actual == pytest.approx(value), name
            else:
                assert actual == value, name


def test_export_merged_sim2_keeps_original_timeseries(workbook, tmp_path):
    paths = [workbook(f"AT00{i}.xlsx", sheets=timestep_sheets(i * 10)) for i in (1, 2)]
    project = Project.from_files(paths, unknown="raise")
    before = [s.timeseries.to_frame() for s in project]
    target = project.to_excel(tmp_path / "annual.xlsx", include_out=True)
    restored = Project.from_excel(target, unknown="raise")
    assert restored.column_names() == ["AT001", "AT002"]
    assert [s.to_var_dict()["QH"] for s in restored] == [10, 20]
    assert all(not s.timeseries.available for s in restored)
    for scenario, frame in zip(project, before):
        pd.testing.assert_frame_equal(scenario.timeseries.to_frame(), frame)


def test_out_is_complete_when_derived_inputs_are_excluded(workbook, tmp_path):
    project = Project.from_excel(workbook(sheets=annual_sheets()))
    full = project.to_excel(tmp_path / "full.xlsx", include_out=True)
    filtered = project.to_excel(
        tmp_path / "filtered.xlsx", include_out=True, include_derived=False,
    )
    pd.testing.assert_frame_equal(
        pd.read_excel(full, sheet_name="OUT"),
        pd.read_excel(filtered, sheet_name="OUT"),
    )


def test_writer_applies_metadata_options_to_both_sheets(workbook, tmp_path):
    project = Project.from_excel(workbook(sheets=annual_sheets()))
    target = write_project_excel(
        project, tmp_path / "custom.xlsx", include_out=True,
        base_columns=["var_name", "default"], include_default=True,
    )
    for sheet in ("IN", "OUT"):
        assert pd.read_excel(target, sheet_name=sheet).columns.tolist() == [
            "var_name", "default", "Demo | B", "Demo | A",
        ]


def test_empty_project_annual_roundtrip(tmp_path):
    target = Project().to_excel(tmp_path / "empty.xlsx", include_out=True)
    assert len(Project.from_excel(target, unknown="raise")) == 0


def test_legacy_values_are_inspectable_but_not_exported(workbook):
    sheets = annual_sheets()
    sheets["IN"].append(["PV_roof_area_percentage", "Old PV", 0.4, 0.6])
    path = workbook(sheets=sheets)
    project = Project.from_excel(path, unknown="ignore")
    assert project[0].legacy == {"PV_roof_area_percentage": 0.4}
    assert project[0].v.PV_roof_share is None
    assert "PV_roof_area_percentage" not in project[0].to_var_dict()
    assert len([w for w in project.warnings if "scenario.legacy" in w]) == 1
    project[0].v.PV_roof_share = project[0].legacy["PV_roof_area_percentage"]
    assert project[0].v.PV_roof_share == 0.4
    assert Project.from_excel(path, unknown="ignore")[0].v.PV_roof_share is None
    with pytest.raises(KeyError):
        Project.from_excel(path, unknown="raise")


def test_conflicting_legacy_values_fail(workbook):
    sheets = annual_sheets()
    sheets["IN"].append(["old_field", "Old", 1, 2])
    sheets["OUT"].append(["old_field", 3, 4])
    with pytest.raises(ValueError, match="Conflicting legacy"):
        Project.from_excel(workbook(sheets=sheets), unknown="ignore")
